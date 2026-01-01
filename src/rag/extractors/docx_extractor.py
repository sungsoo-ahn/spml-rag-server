"""Word document extractor using python-docx."""

from pathlib import Path
from docx import Document

from src.rag.extractors.base import BaseExtractor, ExtractedDocument


class DocxExtractor(BaseExtractor):
    """Extract text from Word documents using python-docx."""

    supported_extensions = [".docx"]

    def extract(self, file_path: Path) -> ExtractedDocument:
        """Extract text content from a Word document."""
        doc = Document(str(file_path))
        paragraphs = []

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                if row_text.strip():
                    paragraphs.append(row_text)

        content = "\n\n".join(paragraphs)

        # Extract metadata
        metadata = {}
        if doc.core_properties:
            props = doc.core_properties
            if props.author:
                metadata["author"] = props.author
            if props.title:
                metadata["title"] = props.title
            if props.subject:
                metadata["subject"] = props.subject

        return ExtractedDocument(
            content=content,
            source_path=str(file_path),
            file_type="docx",
            title=file_path.stem,
            metadata=metadata,
        )
