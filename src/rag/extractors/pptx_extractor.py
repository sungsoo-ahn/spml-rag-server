"""PowerPoint document extractor using python-pptx."""

from pathlib import Path
from pptx import Presentation

from src.rag.extractors.base import BaseExtractor, ExtractedDocument


class PptxExtractor(BaseExtractor):
    """Extract text from PowerPoint presentations using python-pptx."""

    supported_extensions = [".pptx"]

    def extract(self, file_path: Path) -> ExtractedDocument:
        """Extract text content from a PowerPoint presentation."""
        prs = Presentation(str(file_path))
        slides_text = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_content = [f"--- Slide {slide_num} ---"]

            for shape in slide.shapes:
                # Extract text from shapes
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text)

                # Handle tables in slides
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text for cell in row.cells if cell.text.strip()
                        )
                        if row_text.strip():
                            slide_content.append(row_text)

            if len(slide_content) > 1:  # More than just the slide header
                slides_text.append("\n".join(slide_content))

        content = "\n\n".join(slides_text)

        return ExtractedDocument(
            content=content,
            source_path=str(file_path),
            file_type="pptx",
            title=file_path.stem,
            page_count=len(prs.slides),
            metadata={"slide_count": len(prs.slides)},
        )
